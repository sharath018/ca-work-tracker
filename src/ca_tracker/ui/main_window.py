"""Main application window for CA Work Tracker."""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import csv
import logging
from datetime import datetime
from ca_tracker.config import (
    WINDOW_WIDTH, WINDOW_HEIGHT, WINDOW_MIN_WIDTH, WINDOW_MIN_HEIGHT,
    CATEGORIES, STATUS_OPTIONS, BILLABLE_OPTIONS, DATE_FORMAT,
    APP_NAME, APP_VERSION
)
from ca_tracker.database import DatabaseManager
from ca_tracker.utils import get_db_path, get_backup_path, get_logger
from ca_tracker.ui.filters import DateRangeFilter, FilterState
from ca_tracker.ui.reports import MonthlySummaryReport, SummaryReportWindow

logger = get_logger()

try:
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logger.warning("reportlab not installed - PDF export will be unavailable")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not installed - Excel export will be unavailable")


class MainWindow:
    """Main application window for CA Work Tracker."""
    
    def __init__(self, root):
        """Initialize the main window."""
        self.root = root
        self.root.title(f"{APP_NAME} v{APP_VERSION}")
        self.root.geometry(f"{WINDOW_WIDTH}x{WINDOW_HEIGHT}")
        self.root.minsize(WINDOW_MIN_WIDTH, WINDOW_MIN_HEIGHT)
        
        self.selected_id = None
        self.filter_state = FilterState()  # Track current filter state
        self.client_suggestions = []
        
        # Initialize database
        try:
            self.db = DatabaseManager(get_db_path(), get_backup_path())
            self.db.backup_database(max_backups=30)
            logger.info("Database initialized and backed up")
        except Exception as e:
            logger.error(f"Failed to initialize database: {e}")
            messagebox.showerror("Database Error", f"Failed to initialize database: {e}")
            self.root.destroy()
            return
        
        self.create_ui()
        self.load_data()
        self.update_dashboard()
        
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
    
    def create_ui(self):
        """Create the user interface."""
        # Dashboard
        dashboard_frame = tk.LabelFrame(self.root, text="Dashboard", padx=10, pady=10)
        dashboard_frame.pack(fill="x", padx=10, pady=5)
        
        self.total_amount_label = tk.Label(
            dashboard_frame, text="Total Billable Amount: ₹0", font=("Arial", 12, "bold")
        )
        self.total_amount_label.pack(side="left", padx=20)
        
        self.pending_label = tk.Label(
            dashboard_frame, text="Pending Tasks: 0", font=("Arial", 12, "bold")
        )
        self.pending_label.pack(side="left", padx=20)
        
        self.completed_label = tk.Label(
            dashboard_frame, text="Completed Tasks: 0", font=("Arial", 12, "bold")
        )
        self.completed_label.pack(side="left", padx=20)
        
        # Form
        form_frame = tk.LabelFrame(self.root, text="Work Entry", padx=10, pady=10)
        form_frame.pack(fill="x", padx=10, pady=5)
        
        # Row 1
        tk.Label(form_frame, text="Date (DD-MMM-YYYY)").grid(row=0, column=0, sticky="w")
        self.date_entry = tk.Entry(form_frame, width=15)
        self.date_entry.grid(row=1, column=0, padx=5, pady=5)
        self.date_entry.insert(0, datetime.now().strftime(DATE_FORMAT))
        
        tk.Label(form_frame, text="Client").grid(row=0, column=1, sticky="w")
        self.client_combo = ttk.Combobox(form_frame, values=[], width=25)
        self.client_combo.grid(row=1, column=1, padx=5, pady=5)
        self.client_combo.bind('<KeyRelease>', self.update_client_suggestions)

        tk.Label(form_frame, text="Category").grid(row=0, column=2, sticky="w")
        self.category_combo = ttk.Combobox(form_frame, values=CATEGORIES, width=20)
        self.category_combo.grid(row=1, column=2, padx=5, pady=5)
        
        tk.Label(form_frame, text="Hours").grid(row=0, column=3, sticky="w")
        self.hours_entry = tk.Entry(form_frame, width=10)
        self.hours_entry.grid(row=1, column=3, padx=5, pady=5)
        
        tk.Label(form_frame, text="Billable").grid(row=0, column=4, sticky="w")
        self.billable_combo = ttk.Combobox(form_frame, values=BILLABLE_OPTIONS, width=10)
        self.billable_combo.grid(row=1, column=4, padx=5, pady=5)
        
        tk.Label(form_frame, text="Amount (₹)").grid(row=0, column=5, sticky="w")
        self.amount_entry = tk.Entry(form_frame, width=12)
        self.amount_entry.grid(row=1, column=5, padx=5, pady=5)
        
        tk.Label(form_frame, text="Status").grid(row=0, column=6, sticky="w")
        self.status_combo = ttk.Combobox(form_frame, values=STATUS_OPTIONS, width=15)
        self.status_combo.grid(row=1, column=6, padx=5, pady=5)
        
        # Row 2
        tk.Label(form_frame, text="Requested By").grid(row=2, column=0, sticky="w")
        self.requested_by_entry = tk.Entry(form_frame, width=20)
        self.requested_by_entry.grid(row=3, column=0, padx=5, pady=5)
        
        tk.Label(form_frame, text="Summary").grid(row=2, column=1, sticky="w")
        self.summary_entry = tk.Entry(form_frame, width=50)
        self.summary_entry.grid(row=3, column=1, columnspan=3, padx=5, pady=5)
        
        tk.Label(form_frame, text="Notes").grid(row=2, column=4, sticky="w")
        self.notes_entry = tk.Entry(form_frame, width=40)
        self.notes_entry.grid(row=3, column=4, columnspan=3, padx=5, pady=5)
        
        # Buttons
        button_frame = tk.Frame(form_frame)
        button_frame.grid(row=4, column=0, columnspan=7, pady=10)
        
        tk.Button(
            button_frame, text="Add Entry", command=self.add_entry,
            width=15, bg="#4CAF50", fg="white"
        ).pack(side="left", padx=5)
        
        tk.Button(
            button_frame, text="Update Entry", command=self.update_entry,
            width=15, bg="#2196F3", fg="white"
        ).pack(side="left", padx=5)
        
        tk.Button(
            button_frame, text="Clear", command=self.clear_form, width=12
        ).pack(side="left", padx=5)
        
        tk.Button(
            button_frame, text="Delete", command=self.delete_selected,
            width=12, bg="#d9534f", fg="white"
        ).pack(side="left", padx=5)
        
        tk.Button(
            button_frame, text="Export CSV", command=self.export_csv, width=12
        ).pack(side="left", padx=5)

        tk.Button(
            button_frame, text="Export Excel", command=self.export_excel, width=12
        ).pack(side="left", padx=5)
        
        if REPORTLAB_AVAILABLE:
            tk.Button(
                button_frame, text="Export PDF", command=self.export_pdf, width=12
            ).pack(side="left", padx=5)
        
        # Search and Filters
        filter_frame = tk.LabelFrame(self.root, text="Search & Filters", padx=10, pady=10)
        filter_frame.pack(fill="x", padx=10, pady=5)
        
        # Row 1: Search
        search_row = tk.Frame(filter_frame)
        search_row.pack(fill="x", pady=5)
        
        tk.Label(search_row, text="Search").pack(side="left", padx=5)
        self.search_entry = tk.Entry(search_row, width=30)
        self.search_entry.pack(side="left", padx=5)
        
        tk.Button(search_row, text="Search", command=self.search_data).pack(side="left", padx=5)
        tk.Button(search_row, text="Show All", command=self.load_data).pack(side="left", padx=5)
        
        tk.Label(search_row, text="Status Filter").pack(side="left", padx=10)
        self.filter_status = ttk.Combobox(
            search_row, values=[""] + STATUS_OPTIONS, width=15
        )
        self.filter_status.pack(side="left", padx=5)
        
        tk.Button(search_row, text="Apply Filter", command=self.filter_data).pack(side="left", padx=5)
        
        # Row 2: Date Range Filter
        date_filter_row = tk.Frame(filter_frame)
        date_filter_row.pack(fill="x", pady=5)
        
        self.date_range_filter = DateRangeFilter(date_filter_row, callback=self.filter_by_date_range)
        self.date_range_filter.pack()
        
        tk.Button(
            date_filter_row, text="Apply Date Range", 
            command=self.filter_by_date_range, width=15
        ).pack(side="left", padx=5)
        
        tk.Button(
            date_filter_row, text="Monthly Summary", 
            command=self.show_monthly_summary, width=15, bg="#FF9800", fg="white"
        ).pack(side="left", padx=5)
        
        # Table with scrollbars
        table_frame = tk.Frame(self.root)
        table_frame.pack(fill="both", expand=True, padx=10, pady=10)
        
        columns = ("ID", "Date", "Client", "Category", "Summary", "Hours", "Billable", "Amount", "Status", "Requested By", "Notes")
        
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=15)
        
        for col in columns:
            self.tree.heading(col, text=col)
            if col == "Summary":
                self.tree.column(col, width=300)
            elif col == "Notes":
                self.tree.column(col, width=250)
            else:
                self.tree.column(col, width=100)
        
        # Vertical scrollbar
        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        
        # Horizontal scrollbar
        hsb = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(xscrollcommand=hsb.set)
        
        # Grid layout with scrollbars
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)
        
        self.tree.bind("<Double-1>", self.select_record)
    
    def add_entry(self):
        """Add a new work entry."""
        try:
            # Validation
            if not self.client_combo.get().strip():
                messagebox.showwarning("Validation Error", "Client name is required.")
                return
            if not self.category_combo.get().strip():
                messagebox.showwarning("Validation Error", "Please select a Category.")
                return
            if not self.status_combo.get().strip():
                messagebox.showwarning("Validation Error", "Please select a Status.")
                return
            
            # Parse numeric fields
            try:
                hours = float(self.hours_entry.get()) if self.hours_entry.get().strip() else 0.0
                amount = float(self.amount_entry.get()) if self.amount_entry.get().strip() else 0.0
            except ValueError:
                messagebox.showerror("Input Error", "Hours and Amount must be valid numbers.")
                return
            
            # Validate and convert date
            date_str = self.date_entry.get().strip()
            try:
                datetime.strptime(date_str, DATE_FORMAT)
            except ValueError:
                messagebox.showerror("Date Error", f"Date must be in format: {DATE_FORMAT}")
                return
            
            # Convert billable to integer
            billable = 1 if self.billable_combo.get() == "Yes" else 0
            
            self.db.execute("""
                INSERT INTO work_log (
                    date, client, category, summary,
                    hours, billable, amount, status,
                    requested_by, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                date_str,
                self.client_combo.get(),
                self.category_combo.get(),
                self.summary_entry.get(),
                hours,
                billable,
                amount,
                self.status_combo.get(),
                self.requested_by_entry.get(),
                self.notes_entry.get()
            ))
            
            self.db.commit()
            self.load_data()
            self.update_dashboard()
            self.clear_form()
            
            messagebox.showinfo("Success", "Entry added successfully")
            logger.info("New work entry added")
            
        except Exception as e:
            logger.error(f"Error adding entry: {e}")
            messagebox.showerror("Error", f"Failed to add entry: {e}")
        finally:
            self.refresh_client_suggestions()
    
    def load_data(self):
        """Load all work entries into the table."""
        for row in self.tree.get_children():
            self.tree.delete(row)
        
        self.filter_state.clear()
        self.refresh_client_suggestions()
        self.db.execute("SELECT * FROM work_log ORDER BY date DESC, id DESC")
        rows = self.db.fetchall()
        
        for row in rows:
            billable_display = "Yes" if row[6] else "No"
            display_row = (row[0], row[1], row[2], row[3], row[4], row[5], billable_display, row[7], row[8], row[9], row[10])
            self.tree.insert("", tk.END, values=display_row)
    
    def select_record(self, event):
        """Select a record by double-clicking."""
        selected = self.tree.focus()
        
        if not selected:
            return
        
        values = self.tree.item(selected, "values")
        self.selected_id = values[0]
        
        self.date_entry.delete(0, tk.END)
        self.date_entry.insert(0, values[1])
        
        self.client_combo.set(values[2])
        
        self.category_combo.set(values[3])
        
        self.summary_entry.delete(0, tk.END)
        self.summary_entry.insert(0, values[4])
        
        self.hours_entry.delete(0, tk.END)
        self.hours_entry.insert(0, values[5])
        
        self.billable_combo.set(values[6])
        
        self.amount_entry.delete(0, tk.END)
        self.amount_entry.insert(0, values[7])
        
        self.status_combo.set(values[8])
        
        self.requested_by_entry.delete(0, tk.END)
        self.requested_by_entry.insert(0, values[9])
        
        self.notes_entry.delete(0, tk.END)
        self.notes_entry.insert(0, values[10])
        
        self.refresh_client_suggestions()
    
    def update_entry(self):
        """Update the selected work entry."""
        if not self.selected_id:
            messagebox.showwarning("Warning", "Select a record first by double-clicking.")
            return
        
        confirm = messagebox.askyesno(
            "Confirm Update", "Are you sure you want to update this record?"
        )
        
        if not confirm:
            return
        
        try:
            # Parse numeric fields
            try:
                hours = float(self.hours_entry.get()) if self.hours_entry.get().strip() else 0.0
                amount = float(self.amount_entry.get()) if self.amount_entry.get().strip() else 0.0
            except ValueError:
                messagebox.showerror("Input Error", "Hours and Amount must be valid numbers.")
                return
            
            # Validate and convert date
            date_str = self.date_entry.get().strip()
            try:
                datetime.strptime(date_str, DATE_FORMAT)
            except ValueError:
                messagebox.showerror("Date Error", f"Date must be in format: {DATE_FORMAT}")
                return
            
            # Convert billable to integer
            billable = 1 if self.billable_combo.get() == "Yes" else 0
            
            self.db.execute("""
                UPDATE work_log
                SET date=?, client=?, category=?, summary=?,
                    hours=?, billable=?, amount=?, status=?,
                    requested_by=?, notes=?, updated_at=CURRENT_TIMESTAMP
                WHERE id=?
            """, (
                date_str,
                self.client_combo.get(),
                self.category_combo.get(),
                self.summary_entry.get(),
                hours,
                billable,
                amount,
                self.status_combo.get(),
                self.requested_by_entry.get(),
                self.notes_entry.get(),
                self.selected_id
            ))
            
            self.db.commit()
            self.load_data()
            self.update_dashboard()
            self.clear_form()
            
            messagebox.showinfo("Success", "Record updated successfully")
            logger.info(f"Work entry {self.selected_id} updated")
            
        except Exception as e:
            logger.error(f"Error updating entry: {e}")
            messagebox.showerror("Error", f"Failed to update entry: {e}")
        finally:
            self.refresh_client_suggestions()
    
    def search_data(self):
        """Search for work entries."""
        keyword = self.search_entry.get().strip()
        
        if not keyword:
            messagebox.showwarning("Warning", "Please enter a search term.")
            return
        
        for row in self.tree.get_children():
            self.tree.delete(row)
        
        wildcard = f"%{keyword}%"
        
        self.db.execute("""
            SELECT * FROM work_log
            WHERE client LIKE ? OR category LIKE ? OR summary LIKE ?
                  OR status LIKE ? OR requested_by LIKE ?
            ORDER BY date DESC
        """, (wildcard, wildcard, wildcard, wildcard, wildcard))
        
        rows = self.db.fetchall()
        
        for row in rows:
            billable_display = "Yes" if row[6] else "No"
            display_row = (row[0], row[1], row[2], row[3], row[4], row[5], billable_display, row[7], row[8], row[9], row[10])
            self.tree.insert("", tk.END, values=display_row)
        
        # Update filter state
        self.filter_state.set_search(keyword)
    
    def filter_data(self):
        """Filter work entries by status."""
        status = self.filter_status.get()
        
        for row in self.tree.get_children():
            self.tree.delete(row)
        
        if status:
            self.db.execute(
                "SELECT * FROM work_log WHERE status=? ORDER BY date DESC",
                (status,)
            )
            self.filter_state.set_status_filter(status)
        else:
            self.db.execute("SELECT * FROM work_log ORDER BY date DESC")
            self.filter_state.clear()
        
        rows = self.db.fetchall()
        
        for row in rows:
            billable_display = "Yes" if row[6] else "No"
            display_row = (row[0], row[1], row[2], row[3], row[4], row[5], billable_display, row[7], row[8], row[9], row[10])
            self.tree.insert("", tk.END, values=display_row)
    
    def delete_selected(self):
        """Delete the selected work entry."""
        selected = self.tree.selection()
        
        if not selected:
            messagebox.showwarning("Warning", "Please select a record to delete.")
            return
        
        item = self.tree.item(selected[0])
        record_id = item['values'][0]
        
        confirm = messagebox.askyesno("Confirm", "Delete this entry permanently?")
        
        if confirm:
            try:
                self.db.execute("DELETE FROM work_log WHERE id=?", (record_id,))
                self.db.commit()
                self.load_data()
                self.update_dashboard()
                logger.info(f"Work entry {record_id} deleted")
            except Exception as e:
                logger.error(f"Error deleting entry: {e}")
                messagebox.showerror("Error", f"Failed to delete entry: {e}")
    
    def update_dashboard(self):
        """Update the dashboard statistics."""
        try:
            # Total billable amount
            self.db.execute("SELECT SUM(amount) FROM work_log WHERE billable=1")
            total_billable = self.db.fetchone()[0] or 0
            self.total_amount_label.config(text=f"Total Billable Amount: ₹{total_billable:,.2f}")
            
            # Pending tasks
            self.db.execute("SELECT COUNT(*) FROM work_log WHERE status='Pending'")
            pending = self.db.fetchone()[0] or 0
            self.pending_label.config(text=f"Pending Tasks: {pending}")
            
            # Completed tasks
            self.db.execute("SELECT COUNT(*) FROM work_log WHERE status='Completed'")
            completed = self.db.fetchone()[0] or 0
            self.completed_label.config(text=f"Completed Tasks: {completed}")
            
        except Exception as e:
            logger.error(f"Error updating dashboard: {e}")
    
    def export_csv(self):
        """Export work entries to CSV (respects current filter)."""
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV Files", "*.csv")]
        )
        
        if not file_path:
            return
        
        try:
            # Build query based on current filter
            where_clause, params = self.filter_state.get_sql_clause()
            
            if where_clause:
                query = f"SELECT * FROM work_log {where_clause} ORDER BY date DESC"
            else:
                query = "SELECT * FROM work_log ORDER BY date DESC"
            
            self.db.execute(query, params)
            rows = self.db.fetchall()
            
            with open(file_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow([
                    "ID", "Date", "Client", "Category", "Summary",
                    "Hours", "Billable", "Amount", "Status",
                    "Requested By", "Notes", "Created", "Updated"
                ])
                
                for row in rows:
                    billable_display = "Yes" if row[6] else "No"
                    writer.writerow([
                        row[0], row[1], row[2], row[3], row[4],
                        row[5], billable_display, row[7], row[8],
                        row[9], row[10], row[11], row[12]
                    ])
            
            messagebox.showinfo("Success", "CSV exported successfully")
            logger.info(f"CSV export completed: {file_path}")
            
        except Exception as e:
            logger.error(f"Error exporting CSV: {e}")
            messagebox.showerror("Export Error", f"Failed to export CSV: {e}")
    
    def export_excel(self):
        """Export work entries to Excel .xlsx (respects current filter)."""
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Error", "openpyxl is not installed.\nRun: pip install openpyxl")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")]
        )
        
        if not file_path:
            return
        
        try:
            where_clause, params = self.filter_state.get_sql_clause()
            query = f"SELECT * FROM work_log {where_clause} ORDER BY date DESC" if where_clause else "SELECT * FROM work_log ORDER BY date DESC"
            self.db.execute(query, params)
            rows = self.db.fetchall()
            
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Work Log"
            
            headers = [
                "ID", "Date", "Client", "Category", "Summary",
                "Hours", "Billable", "Amount", "Status",
                "Requested By", "Notes", "Created", "Updated"
            ]
            sheet.append(headers)
            
            for row in rows:
                billable_display = "Yes" if row[6] else "No"
                sheet.append([
                    row[0], row[1], row[2], row[3], row[4],
                    row[5], billable_display, row[7], row[8],
                    row[9], row[10], row[11], row[12]
                ])
            
            bold_font = Font(bold=True)
            for cell in sheet[1]:
                cell.font = bold_font
            
            column_widths = [12, 15, 25, 20, 40, 10, 12, 15, 15, 25, 35, 20, 20]
            for i, width in enumerate(column_widths, start=1):
                sheet.column_dimensions[get_column_letter(i)].width = width
            
            workbook.save(file_path)
            messagebox.showinfo("Success", "Excel exported successfully")
            logger.info(f"Excel export completed: {file_path}")
        except Exception as e:
            logger.error(f"Error exporting Excel: {e}")
            messagebox.showerror("Export Error", f"Failed to export Excel: {e}")
        
    def export_pdf(self):
        """Export work entries to PDF (respects current filter)."""
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Error", "reportlab is not installed.\nRun: pip install reportlab")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF Files", "*.pdf")]
        )
        
        if not file_path:
            return
        
        try:
            doc = SimpleDocTemplate(file_path)
            elements = []
            styles = getSampleStyleSheet()
            
            title = Paragraph(f"CA Work Tracker Report - {datetime.now().strftime('%d-%b-%Y')}", styles['Title'])
            elements.append(title)
            elements.append(Spacer(1, 12))
            
            # Build query based on current filter
            where_clause, params = self.filter_state.get_sql_clause()
            
            if where_clause:
                query = f"SELECT * FROM work_log {where_clause} ORDER BY date DESC"
            else:
                query = "SELECT * FROM work_log ORDER BY date DESC"
            
            self.db.execute(query, params)
            rows = self.db.fetchall()
            
            data = [[
                "Date", "Client", "Category", "Summary",
                "Hours", "Amount", "Status", "Billable"
            ]]
            
            for row in rows:
                billable_display = "Yes" if row[6] else "No"
                data.append([
                    str(row[1]), str(row[2]), str(row[3]),
                    str(row[4])[:30], str(row[5]), f"₹{row[7]:.2f}",
                    str(row[8]), billable_display
                ])
            
            table = Table(data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8)
            ]))
            
            elements.append(table)
            doc.build(elements)
            
            messagebox.showinfo("Success", "PDF exported successfully")
            logger.info(f"PDF export completed: {file_path}")
            
        except Exception as e:
            logger.error(f"Error exporting PDF: {e}")
            messagebox.showerror("Export Error", f"Failed to export PDF: {e}")
    
    def clear_form(self):
        """Clear all form fields."""
        self.selected_id = None
        
        self.date_entry.delete(0, tk.END)
        self.date_entry.insert(0, datetime.now().strftime(DATE_FORMAT))
        
        self.client_combo.set("")
        self.category_combo.set("")
        self.summary_entry.delete(0, tk.END)
        self.hours_entry.delete(0, tk.END)
        self.billable_combo.set("")
        self.amount_entry.delete(0, tk.END)
        self.status_combo.set("")
        self.requested_by_entry.delete(0, tk.END)
        self.notes_entry.delete(0, tk.END)
    
    def get_client_list(self):
        """Return distinct client names for autocomplete suggestions."""
        try:
            self.db.execute(
                "SELECT DISTINCT client FROM work_log WHERE client IS NOT NULL AND client != '' ORDER BY client COLLATE NOCASE ASC"
            )
            rows = self.db.fetchall()
            return [row[0] for row in rows if row[0]]
        except Exception as e:
            logger.warning(f"Unable to refresh client suggestions: {e}")
            return []

    def refresh_client_suggestions(self):
        """Refresh the client autocomplete values."""
        self.client_suggestions = self.get_client_list()
        self.client_combo['values'] = self.client_suggestions

    def update_client_suggestions(self, event=None):
        """Filter autocomplete options as the user types."""
        typed = self.client_combo.get().strip().lower()
        if typed and self.client_suggestions:
            filtered = [client for client in self.client_suggestions if typed in client.lower()]
            self.client_combo['values'] = filtered
        else:
            self.client_combo['values'] = self.client_suggestions

    def filter_by_date_range(self):
        """Filter work entries by date range."""
        from_date, to_date = self.date_range_filter.get_range()
        
        if not from_date or not to_date:
            messagebox.showerror("Error", "Invalid date format. Use DD-MMM-YYYY")
            return
        
        if from_date > to_date:
            messagebox.showerror("Error", "From date must be before To date")
            return
        
        for row in self.tree.get_children():
            self.tree.delete(row)
        
        from_str = from_date.strftime(DATE_FORMAT)
        to_str = to_date.strftime(DATE_FORMAT)
        
        self.db.execute(
            "SELECT * FROM work_log WHERE date BETWEEN ? AND ? ORDER BY date DESC",
            (from_str, to_str)
        )
        
        rows = self.db.fetchall()
        
        for row in rows:
            billable_display = "Yes" if row[6] else "No"
            display_row = (row[0], row[1], row[2], row[3], row[4], row[5], billable_display, row[7], row[8], row[9], row[10])
            self.tree.insert("", tk.END, values=display_row)
        
        # Update filter state
        self.filter_state.set_date_range(from_date, to_date)
        
        logger.info(f"Filtered by date range: {from_str} to {to_str}")
    
    def show_monthly_summary(self):
        """Show monthly summary report."""
        from_date, to_date = self.date_range_filter.get_range()
        
        if not from_date or not to_date:
            messagebox.showerror("Error", "Invalid date format. Use DD-MMM-YYYY")
            return
        
        if from_date > to_date:
            messagebox.showerror("Error", "From date must be before To date")
            return
        
        try:
            report = MonthlySummaryReport(self.db, from_date, to_date)
            SummaryReportWindow(self.root, report)
            logger.info(f"Generated monthly summary for {from_date.strftime(DATE_FORMAT)} to {to_date.strftime(DATE_FORMAT)}")
        except Exception as e:
            logger.error(f"Error generating summary report: {e}")
            messagebox.showerror("Error", f"Failed to generate report: {e}")
    
    def on_closing(self):
        """Handle application closing."""
        try:
            self.db.close()
            logger.info("Application closed")
        except Exception as e:
            logger.error(f"Error closing application: {e}")
        finally:
            self.root.destroy()
